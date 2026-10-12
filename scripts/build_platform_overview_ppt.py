from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports"
OUT_PATH = OUT_DIR / "智能仓储AGV调度仿真优化平台-5页说明.pptx"

WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

COLORS = {
    "bg": RGBColor(248, 250, 252),
    "ink": RGBColor(31, 41, 55),
    "muted": RGBColor(88, 98, 112),
    "teal": RGBColor(0, 109, 119),
    "orange": RGBColor(249, 115, 22),
    "blue": RGBColor(37, 99, 235),
    "green": RGBColor(22, 163, 74),
    "line": RGBColor(210, 218, 228),
    "white": RGBColor(255, 255, 255),
    "dark": RGBColor(15, 23, 42),
}


def set_text(shape, text, size=18, color="ink", bold=False, align=None):
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    run = p.runs[0]
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = COLORS[color]


def add_textbox(slide, left, top, width, height, text, size=18, color="ink", bold=False, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    set_text(box, text, size=size, color=color, bold=bold, align=align)
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, Inches(0.62), Inches(0.42), Inches(9.8), Inches(0.55), title, 26, "dark", True)
    if subtitle:
        add_textbox(slide, Inches(0.66), Inches(0.96), Inches(10.8), Inches(0.38), subtitle, 13, "muted")
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(1.34), Inches(1.2), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["orange"]
    bar.line.fill.background()


def add_card(slide, left, top, width, height, title, body, accent="teal"):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["white"]
    shape.line.color.rgb = COLORS["line"]
    shape.line.width = Pt(1)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Inches(0.08), height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = COLORS[accent]
    stripe.line.fill.background()
    add_textbox(slide, left + Inches(0.22), top + Inches(0.16), width - Inches(0.34), Inches(0.32), title, 15, "dark", True)
    box = add_textbox(slide, left + Inches(0.22), top + Inches(0.58), width - Inches(0.34), height - Inches(0.7), body, 12.5, "muted")
    box.text_frame.paragraphs[0].line_spacing = 1.1
    return shape


def add_bullets(slide, left, top, width, height, items, size=15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = COLORS["ink"]
        p.space_after = Pt(8)
    return box


def add_footer(slide, page_no):
    add_textbox(slide, Inches(11.6), Inches(7.0), Inches(1.1), Inches(0.22), f"0{page_no}", 10, "muted", True, PP_ALIGN.RIGHT)


def set_background(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = COLORS["bg"]


def build_deck():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    blank = prs.slide_layouts[6]

    # 1
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "平台整体流程", "从仓库拓扑建模到算法评估，形成可验证的 AGV 调度闭环")
    steps = [
        ("拓扑建模", "绘制道路、障碍、库位、站台、输送带和 AGV 节点"),
        ("场景配置", "绑定拓扑，配置机器人、货物、任务和运行策略"),
        ("仿真运行", "启动调度过程，实时观察车辆路径和任务状态"),
        ("指标评估", "统计路径长度、规划耗时、重规划次数和冲突情况"),
    ]
    left = Inches(0.82)
    for idx, (title, body) in enumerate(steps):
        x = left + Inches(idx * 3.05)
        add_card(slide, x, Inches(2.0), Inches(2.45), Inches(1.62), title, body, ["teal", "blue", "orange", "green"][idx])
        if idx < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, x + Inches(2.52), Inches(2.53), Inches(0.38), Inches(0.34))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COLORS["line"]
            arrow.line.fill.background()
    add_bullets(
        slide,
        Inches(0.95),
        Inches(4.25),
        Inches(11.4),
        Inches(1.6),
        [
            "平台把原本分散的地图、设备、货物和任务数据放到同一套仿真流程中。",
            "拓扑和场景先校验再运行，减少因为数据不一致导致的仿真失败。",
            "算法输出通过统一指标比较，便于答辩展示和后续调度策略优化。",
        ],
        15,
    )
    add_footer(slide, 1)

    # 2
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "拓扑编辑器解决的问题", "把仓库布局转成算法可以理解、仿真可以运行的结构化地图")
    add_card(
        slide,
        Inches(0.75),
        Inches(1.8),
        Inches(3.6),
        Inches(3.95),
        "建模问题",
        "真实仓库布局包含道路、库位、站台、输送带和障碍区域。人工描述难以直接用于路径规划，修改后也难以快速验证。",
        "teal",
    )
    add_card(
        slide,
        Inches(4.85),
        Inches(1.8),
        Inches(3.6),
        Inches(3.95),
        "校验问题",
        "节点越界、节点重叠、设备放在障碍上、输送带断开等问题会导致仿真无法启动。编辑器在保存前给出明确检查结果。",
        "orange",
    )
    add_card(
        slide,
        Inches(8.95),
        Inches(1.8),
        Inches(3.6),
        Inches(3.95),
        "复用问题",
        "同一张仓库地图可以绑定多个场景，用于测试不同任务量、机器人数量和调度策略，减少重复建模成本。",
        "blue",
    )
    add_footer(slide, 2)

    # 3
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "场景配置解决的问题", "把静态仓库地图变成一次可执行、可复现的业务仿真")
    add_bullets(
        slide,
        Inches(0.82),
        Inches(1.75),
        Inches(5.6),
        Inches(4.7),
        [
            "场景明确本次仿真使用哪一张拓扑，避免任务数据和地图数据脱节。",
            "机器人参数与拓扑设备节点绑定，车辆起点、类型和可服务区域保持一致。",
            "货物与任务被统一管理，系统可以判断哪些货物待分配、执行中或已完成。",
            "调度策略和路径策略在场景中配置，方便在相同条件下对比算法效果。",
            "拓扑发生变化后，场景同步更新相关设备和节点，降低引用失效风险。",
        ],
        15,
    )
    matrix = [
        ("输入", "拓扑 ID、机器人、货物、任务、策略"),
        ("处理", "同步拓扑设备，检查引用和参数"),
        ("输出", "可用于仿真的完整场景数据"),
    ]
    for idx, (head, body) in enumerate(matrix):
        add_card(slide, Inches(7.0), Inches(1.75 + idx * 1.45), Inches(5.2), Inches(1.05), head, body, ["teal", "orange", "green"][idx])
    add_footer(slide, 3)

    # 4
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "场景接口约束", "接口约束保证拓扑、场景和仿真运行使用同一套有效数据")
    rows = [
        ("拓扑引用", "topologyId 必须存在，场景只能绑定已经保存的拓扑"),
        ("设备一致", "场景机器人必须能映射到拓扑设备节点，类型不能错配"),
        ("节点有效", "任务涉及的库位、站台、入口和起点必须来自当前拓扑"),
        ("同步更新", "拓扑修改后，场景需要同步设备、节点和任务引用"),
        ("运行校验", "启动仿真前检查拓扑和场景一致性，失败时返回具体错误"),
    ]
    top = Inches(1.65)
    for i, (constraint, detail) in enumerate(rows):
        y = top + Inches(i * 0.86)
        color = "teal" if i % 2 == 0 else "blue"
        pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.82), y, Inches(2.0), Inches(0.5))
        pill.fill.solid()
        pill.fill.fore_color.rgb = COLORS[color]
        pill.line.fill.background()
        add_textbox(slide, Inches(0.93), y + Inches(0.09), Inches(1.78), Inches(0.25), constraint, 13, "white", True, PP_ALIGN.CENTER)
        add_textbox(slide, Inches(3.05), y + Inches(0.03), Inches(9.0), Inches(0.48), detail, 15, "ink")
    add_card(
        slide,
        Inches(0.82),
        Inches(6.08),
        Inches(11.45),
        Inches(0.62),
        "接口设计重点",
        "先保存、再校验、后运行。任何缺失引用或类型错配都会在进入仿真前被拦截。",
        "orange",
    )
    add_footer(slide, 4)

    # 5
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "任务分配约束", "任务分配关注谁能接单、能否到达、执行成本是否更低")
    add_card(
        slide,
        Inches(0.78),
        Inches(1.65),
        Inches(3.75),
        Inches(4.7),
        "可分配对象",
        "只有空闲机器人可以接收新任务。同一机器人不能同时执行多个任务，同一货物不能被重复派单。",
        "teal",
    )
    add_card(
        slide,
        Inches(4.78),
        Inches(1.65),
        Inches(3.75),
        Inches(4.7),
        "业务匹配",
        "任务必须绑定有效货物。货物所在库位、服务区域和机器人类型需要匹配，例如 CTU 任务优先由 CTU 设备处理。",
        "orange",
    )
    add_card(
        slide,
        Inches(8.78),
        Inches(1.65),
        Inches(3.75),
        Inches(4.7),
        "路径可达",
        "机器人必须能从当前位置到达取货点，并继续到达交接点或目标站台。不可达任务不会进入正常调度。",
        "green",
    )
    add_textbox(
        slide,
        Inches(0.95),
        Inches(6.55),
        Inches(11.2),
        Inches(0.35),
        "分配策略综合路径距离、设备状态、服务区域和任务优先级，选择成本更低的执行方案。",
        15,
        "dark",
        True,
        PP_ALIGN.CENTER,
    )
    add_footer(slide, 5)

    return prs


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    deck = build_deck()
    deck.save(OUT_PATH)
    print(OUT_PATH)
