from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


ROOT = Path(r"D:\Users\pyc\Desktop\智能仓储AGV调度仿真优化平台\仿真平台")
OUTPUT_PATH = ROOT / "shared" / "智能仓储AGV调度仿真优化平台-学校答辩版-10页.pptx"
NOTES_PATH = ROOT / "shared" / "智能仓储AGV调度仿真优化平台-学校答辩版-10页-讲稿.md"
HERO_IMAGE = ROOT / "frontend" / "src" / "assets" / "hero.png"


BG = RGBColor(248, 250, 252)
CARD = RGBColor(255, 255, 255)
PRIMARY = RGBColor(51, 65, 85)
SECONDARY = RGBColor(100, 116, 139)
ACCENT = RGBColor(249, 115, 22)
ACCENT_LIGHT = RGBColor(255, 237, 213)
BORDER = RGBColor(226, 232, 240)
SUCCESS = RGBColor(22, 163, 74)
WARNING = RGBColor(220, 38, 38)
INFO = RGBColor(14, 116, 144)


def set_slide_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    top_bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Cm(0.45))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = ACCENT
    top_bar.line.fill.background()


def set_font(run, size: int, bold: bool = False, color: RGBColor = PRIMARY, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=PRIMARY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, items, size=18, color=PRIMARY, level0_bold=False, bullet_color=ACCENT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(6)
        p.bullet = True
        p.font.color.rgb = bullet_color
        p.text = ""
        run = p.add_run()
        run.text = item
        set_font(run, size=size, bold=level0_bold, color=color)
    return box


def add_card(slide, left, top, width, height, title, body, accent=ACCENT, title_size=18, body_size=13):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(1)
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Cm(0.22), height)
    tag.fill.solid()
    tag.fill.fore_color.rgb = accent
    tag.line.fill.background()
    add_textbox(slide, left + Cm(0.45), top + Cm(0.28), width - Cm(0.7), Cm(0.75), title, size=title_size, bold=True)
    add_textbox(slide, left + Cm(0.45), top + Cm(1.1), width - Cm(0.7), height - Cm(1.3), body, size=body_size, color=SECONDARY)
    return shape


def add_title_block(slide, title: str, subtitle: str) -> None:
    add_textbox(slide, Cm(1.2), Cm(0.9), Cm(22), Cm(1), title, size=26, bold=True)
    add_textbox(slide, Cm(1.2), Cm(2.0), Cm(22), Cm(0.8), subtitle, size=12, color=SECONDARY)


def add_process(slide, left, top, labels, width=Cm(4.0), height=Cm(1.4), gap=Cm(0.25)):
    x = left
    for index, label in enumerate(labels):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, x, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = ACCENT if index in {0, len(labels) - 1} else ACCENT_LIGHT
        shape.line.color.rgb = ACCENT
        text_frame = shape.text_frame
        text_frame.clear()
        p = text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        set_font(run, size=12, bold=True, color=CARD if index in {0, len(labels) - 1} else PRIMARY)
        x += width + gap


def add_footer(slide, text: str) -> None:
    add_textbox(slide, Cm(1.2), Cm(18.0), Cm(31.0), Cm(0.5), text, size=10, color=SECONDARY)


def build_presentation() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Cm(33.867)
    presentation.slide_height = Cm(19.05)
    return presentation


def slide_cover(slide) -> None:
    set_slide_bg(slide)
    overlay = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Cm(0.9), Cm(1.2), Cm(18.2), Cm(15.7))
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = CARD
    overlay.line.color.rgb = BORDER

    add_textbox(slide, Cm(1.6), Cm(2.0), Cm(15.6), Cm(2.3), "智能仓储 AGV 调度仿真优化平台\n设计与实现", size=26, bold=True)
    add_textbox(slide, Cm(1.6), Cm(5.0), Cm(15.0), Cm(1.8), "面向 CTU、四向车与输送线协同调度的仿真优化平台", size=16, color=SECONDARY)
    add_card(
        slide,
        Cm(1.6),
        Cm(7.0),
        Cm(15.2),
        Cm(3.2),
        "项目定位",
        "本项目不是单一算法程序，而是一套从拓扑建模、场景配置、调度仿真、算法对比到结果回放的完整平台。",
        accent=ACCENT,
        title_size=17,
        body_size=14,
    )
    add_textbox(slide, Cm(1.6), Cm(11.0), Cm(10), Cm(0.8), "姓名：__________", size=14, color=SECONDARY)
    add_textbox(slide, Cm(1.6), Cm(11.9), Cm(10), Cm(0.8), "学院：__________", size=14, color=SECONDARY)
    add_textbox(slide, Cm(1.6), Cm(12.8), Cm(10), Cm(0.8), "指导教师：__________", size=14, color=SECONDARY)
    add_textbox(slide, Cm(1.6), Cm(13.7), Cm(10), Cm(0.8), "答辩日期：2026 年 7 月", size=14, color=SECONDARY)

    if HERO_IMAGE.exists():
        slide.shapes.add_picture(str(HERO_IMAGE), Cm(19.8), Cm(2.0), width=Cm(12.8))
    badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Cm(20.4), Cm(13.0), Cm(11.3), Cm(2.6))
    badge.fill.solid()
    badge.fill.fore_color.rgb = PRIMARY
    badge.line.fill.background()
    add_textbox(slide, Cm(21.0), Cm(13.45), Cm(10.2), Cm(1.5), "关键词：网格拓扑 / 协同调度 / 实时仿真 / 算法对比", size=13, color=CARD)


def slide_background(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "项目背景与问题定义", "先说明为什么要做这个平台，再说明平台要解决什么问题。")
    add_card(slide, Cm(1.2), Cm(3.2), Cm(7.5), Cm(3.0), "设备类型多", "CTU、四向车、输送线、缓存位和站台同时参与作业，调度关系复杂。", accent=ACCENT)
    add_card(slide, Cm(9.2), Cm(3.2), Cm(7.5), Cm(3.0), "路径冲突频繁", "多设备共享通道与交接点，容易出现等待、阻塞和死锁风险。", accent=INFO)
    add_card(slide, Cm(17.2), Cm(3.2), Cm(7.5), Cm(3.0), "资源瓶颈明显", "输送线、缓存位和站台的容量有限，会形成系统吞吐瓶颈。", accent=WARNING)
    add_card(slide, Cm(25.2), Cm(3.2), Cm(7.5), Cm(3.0), "真实试错代价高", "直接在真实仓库中修改调度策略风险高、验证慢、成本大。", accent=SUCCESS)
    add_card(
        slide,
        Cm(1.2),
        Cm(7.0),
        Cm(31.8),
        Cm(4.4),
        "传统方法的不足",
        "很多研究只关注静态最短路径或单设备调度，缺少可视化、可回放、可比较的实验环境，难以支撑仓储系统级验证。",
        accent=PRIMARY,
        title_size=18,
        body_size=15,
    )
    add_card(
        slide,
        Cm(1.2),
        Cm(12.1),
        Cm(31.8),
        Cm(4.0),
        "课题目标",
        "构建一个可建模、可运行、可优化、可分析的智能仓储 AGV 调度仿真平台，为后续算法研究和真实系统落地提供试验基础。",
        accent=ACCENT,
        title_size=18,
        body_size=15,
    )
    add_footer(slide, "本页回答两个问题：为什么要做这个平台，以及平台要解决哪些系统级痛点。")


def slide_goals(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "平台目标与总体闭环", "平台的价值不在单点功能，而在从建模到分析的完整闭环。")
    add_card(slide, Cm(1.2), Cm(3.3), Cm(7.6), Cm(2.8), "目标 1", "支持仓储拓扑建模，准确表达道路、障碍、输送带和设备。", accent=ACCENT)
    add_card(slide, Cm(9.2), Cm(3.3), Cm(7.6), Cm(2.8), "目标 2", "支持 CTU、四向车与输送线协同调度仿真。", accent=INFO)
    add_card(slide, Cm(17.2), Cm(3.3), Cm(7.6), Cm(2.8), "目标 3", "支持多策略算法实验，实现任务分配与路径策略对比。", accent=SUCCESS)
    add_card(slide, Cm(25.2), Cm(3.3), Cm(7.6), Cm(2.8), "目标 4", "支持运行回放、异常解释和结果分析。", accent=WARNING)
    add_textbox(slide, Cm(1.2), Cm(7.1), Cm(14), Cm(0.7), "平台使用闭环", size=18, bold=True)
    add_process(
        slide,
        Cm(1.2),
        Cm(8.1),
        ["建拓扑", "配场景", "跑仿真", "看指标", "做对比", "回放分析"],
        width=Cm(4.9),
        height=Cm(1.6),
        gap=Cm(0.15),
    )
    add_card(
        slide,
        Cm(1.2),
        Cm(11.4),
        Cm(31.8),
        Cm(4.4),
        "核心结论",
        "本平台不是只做“会动的动画界面”，而是面向仓储调度研究的实验平台：输入是拓扑和参数，输出是实时状态、运行指标、异常解释和策略对比结果。",
        accent=PRIMARY,
        title_size=18,
        body_size=15,
    )
    add_footer(slide, "答辩时要强调：平台的重点是形成建模、运行、优化、分析的完整闭环。")


def slide_architecture(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "系统总体架构", "平台采用前后端分层架构，界面层、服务层和核心层职责清晰。")

    add_card(slide, Cm(1.2), Cm(3.4), Cm(31.8), Cm(3.0), "表现层", "总控驾驶舱、拓扑编辑器、场景配置台、仿真工作台、算法实验室、回放与报告", accent=ACCENT, title_size=20, body_size=15)
    add_card(slide, Cm(1.2), Cm(7.1), Cm(31.8), Cm(3.0), "服务层", "拓扑管理、场景管理、仿真控制、算法评估、运行记录管理、事件推送", accent=INFO, title_size=20, body_size=15)
    add_card(slide, Cm(1.2), Cm(10.8), Cm(31.8), Cm(3.0), "核心层", "路径规划、任务分配、输送线资源控制、缓存位/站台约束、死锁检测、指标统计", accent=SUCCESS, title_size=20, body_size=15)

    add_card(
        slide,
        Cm(1.2),
        Cm(14.4),
        Cm(15.4),
        Cm(2.5),
        "技术栈",
        "前端：React + Vite + TypeScript\n后端：FastAPI\n通信：WebSocket\n数据：SQLite + JSON 样例",
        accent=PRIMARY,
        title_size=17,
        body_size=13,
    )
    add_card(
        slide,
        Cm(17.6),
        Cm(14.4),
        Cm(15.4),
        Cm(2.5),
        "设计思想",
        "将页面交互、仿真控制和调度逻辑解耦，便于后续扩展真实数据接入、多层仿真和新算法。",
        accent=WARNING,
        title_size=17,
        body_size=13,
    )
    add_footer(slide, "这一页要告诉老师：这是一个平台型系统，而不是单一算法脚本。")


def slide_modules(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "六大功能模块总览", "每个页面各司其职，共同支撑智能仓储调度仿真闭环。")
    modules = [
        ("总控驾驶舱", "查看平台全局状态、运行记录和关键指标。"),
        ("拓扑编辑器", "搭建仓储空间模型，放置道路、设备、站台和资源。"),
        ("场景配置台", "配置设备参数、订单生成规则、调度策略和故障注入。"),
        ("仿真工作台", "实时查看设备运动、任务阶段、资源占用和异常告警。"),
        ("算法实验室", "对比任务分配、路径规划和输送线策略的运行表现。"),
        ("回放与报告", "基于事件和指标对一次仿真进行复盘和结果分析。"),
    ]
    positions = [
        (Cm(1.2), Cm(3.5)),
        (Cm(12.0), Cm(3.5)),
        (Cm(22.8), Cm(3.5)),
        (Cm(1.2), Cm(9.8)),
        (Cm(12.0), Cm(9.8)),
        (Cm(22.8), Cm(9.8)),
    ]
    accents = [ACCENT, INFO, SUCCESS, WARNING, PRIMARY, ACCENT]
    for (title, body), (left, top), accent in zip(modules, positions, accents):
        add_card(slide, left, top, Cm(9.7), Cm(4.7), title, body, accent=accent, title_size=18, body_size=14)
    add_footer(slide, "这六个模块分别对应平台的入口、建模、配置、运行、实验和分析阶段。")


def slide_topology_scenario(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "拓扑编辑器与场景配置台", "前者负责空间结构，后者负责运行参数，两者分工明确。")
    add_card(
        slide,
        Cm(1.2),
        Cm(3.2),
        Cm(15.2),
        Cm(5.8),
        "网格化拓扑模型",
        "采用更贴近仓内行走图的网格建模方式，能更直观表达道路、障碍、输送带和多设备活动空间。",
        accent=ACCENT,
        title_size=19,
        body_size=14,
    )
    add_card(slide, Cm(1.8), Cm(6.6), Cm(4.4), Cm(2.0), "格子层", "road / blocked / conveyor", accent=INFO, title_size=15, body_size=12)
    add_card(slide, Cm(6.6), Cm(6.6), Cm(4.4), Cm(2.0), "节点层", "storage / station / buffer / entry / gate / robot / lift", accent=SUCCESS, title_size=15, body_size=11)
    add_card(slide, Cm(11.4), Cm(6.6), Cm(4.4), Cm(2.0), "资源层", "conveyor / buffer_slot / station / lift", accent=WARNING, title_size=15, body_size=11)

    add_card(
        slide,
        Cm(17.2),
        Cm(3.2),
        Cm(15.8),
        Cm(5.8),
        "场景配置职责",
        "在不改变空间结构的前提下，配置速度、电量、启用状态、订单速率、调度策略、路径策略、输送线阈值和故障注入。",
        accent=PRIMARY,
        title_size=19,
        body_size=14,
    )
    add_bullets(
        slide,
        Cm(17.8),
        Cm(6.1),
        Cm(14.2),
        Cm(2.4),
        [
            "设备参数：速度 / 电量 / 启用状态",
            "订单参数：速率 / 数量 / 插单 Tick",
            "策略参数：Greedy / Hungarian、A* / Dijkstra",
            "故障注入：设备故障、输送线拥堵",
        ],
        size=13,
        color=PRIMARY,
    )
    add_card(
        slide,
        Cm(1.2),
        Cm(10.1),
        Cm(31.8),
        Cm(5.3),
        "关键设计点",
        "CTU 和四向车不在场景页手工创建，而是直接来自拓扑中的设备节点。这样可以保证设备身份、起点位置和仿真对象保持一致，避免拓扑与场景脱节。",
        accent=ACCENT,
        title_size=18,
        body_size=15,
    )
    add_footer(slide, "这一页要让老师理解：拓扑建模和场景配置是解耦的，但又通过设备节点强关联。")


def slide_simulation(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "仿真工作台与协同调度闭环", "平台模拟的不是简单点到点移动，而是完整的任务阶段流转。")
    add_process(
        slide,
        Cm(1.2),
        Cm(3.4),
        ["CTU 取货", "四向车取货", "交接到输送线", "等待输送线", "输送中", "站台排队", "站台交付"],
        width=Cm(4.5),
        height=Cm(1.45),
        gap=Cm(0.08),
    )
    add_card(slide, Cm(1.2), Cm(6.0), Cm(7.6), Cm(4.1), "场景棋盘", "统一展示网格、业务节点、资源对象和运行中的设备位置。", accent=ACCENT, title_size=18, body_size=14)
    add_card(slide, Cm(9.2), Cm(6.0), Cm(7.6), Cm(4.1), "设备状态", "跟踪 CTU 和四向车的当前位置、目标格、状态、等待 Tick 和当前任务。", accent=INFO, title_size=18, body_size=14)
    add_card(slide, Cm(17.2), Cm(6.0), Cm(7.6), Cm(4.1), "任务阶段", "观察直送、交接、上输送线、缓存等待和站台交付的完整闭环。", accent=SUCCESS, title_size=18, body_size=14)
    add_card(slide, Cm(25.2), Cm(6.0), Cm(7.8), Cm(4.1), "资源占用", "单独观察输送线、缓存位、站台队列以及全局等待情况。", accent=WARNING, title_size=18, body_size=14)
    add_card(
        slide,
        Cm(1.2),
        Cm(11.2),
        Cm(31.8),
        Cm(4.6),
        "平台价值",
        "仿真工作台把空间结构、设备状态、任务过程和资源约束同时放到一个视角中，让调度过程既能“跑起来”，也能“看得懂”。",
        accent=PRIMARY,
        title_size=18,
        body_size=15,
    )
    add_footer(slide, "答辩时建议配一张运行截图，这一页最能体现平台不是普通演示动画。")


def slide_algorithms(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "后端调度与算法机制", "平台的核心能力来自路径、任务和资源三类逻辑的耦合。")
    add_card(slide, Cm(1.2), Cm(3.3), Cm(10.0), Cm(5.4), "路径规划", "A* 与 Dijkstra\n网格四邻域搜索\n障碍绕行\n输送带低成本偏好\n提升机跨层连通", accent=ACCENT, title_size=18, body_size=14)
    add_card(slide, Cm(11.7), Cm(3.3), Cm(10.0), Cm(5.4), "任务分配", "Greedy\nHungarian\n根据设备类型、当前位置和任务模式做分配\n支持 CTU 直送与跨区输送线协同", accent=INFO, title_size=18, body_size=14)
    add_card(slide, Cm(22.2), Cm(3.3), Cm(10.8), Cm(5.4), "资源控制", "输送线容量阈值\n缓存位占用\n站台排队\n拥堵检测\n死锁告警与建议", accent=SUCCESS, title_size=18, body_size=14)
    add_card(
        slide,
        Cm(1.2),
        Cm(9.4),
        Cm(15.4),
        Cm(5.2),
        "实时事件流",
        "后端运行仿真，通过 WebSocket 持续推送 tick、robot_position、task_state_changed、resource_occupied、alert、metrics_snapshot 等事件。",
        accent=PRIMARY,
        title_size=18,
        body_size=14,
    )
    add_card(
        slide,
        Cm(17.2),
        Cm(9.4),
        Cm(15.8),
        Cm(5.2),
        "关键结论",
        "平台不是只求一条最短路，而是把空间路径、任务分配和资源约束一起纳入调度逻辑，因此更接近真实仓储协同过程。",
        accent=WARNING,
        title_size=18,
        body_size=14,
    )
    add_footer(slide, "这一页主要回答“平台内部怎么跑起来”，是技术实现的核心。")


def slide_results(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "算法实验、结果分析与演示流程", "平台既能做策略对比，也能对运行结果进行回放和分析。")
    add_card(slide, Cm(1.2), Cm(3.2), Cm(15.2), Cm(5.4), "算法实验室", "可比较三类策略：任务分配、路径规划、输送线放行。\n核心指标：吞吐量、平均完成时长、输送线利用率、拥堵次数、死锁次数、紧急插单完成数。", accent=ACCENT, title_size=19, body_size=14)
    add_card(slide, Cm(17.2), Cm(3.2), Cm(15.8), Cm(5.4), "回放与报告", "可查看指标摘要、输送线利用率趋势、事件热力图和异常时间轴，用于复盘一次完整运行并定位问题。", accent=INFO, title_size=19, body_size=14)
    add_textbox(slide, Cm(1.2), Cm(9.4), Cm(15), Cm(0.7), "现场演示顺序", size=18, bold=True)
    add_process(
        slide,
        Cm(1.2),
        Cm(10.3),
        ["驾驶舱", "拓扑", "场景", "仿真", "实验室", "回放报告"],
        width=Cm(5.1),
        height=Cm(1.5),
        gap=Cm(0.1),
    )
    add_card(
        slide,
        Cm(1.2),
        Cm(12.6),
        Cm(31.8),
        Cm(3.2),
        "本页结论",
        "前面的模块负责把系统“跑起来”，算法实验室负责“比出差异”，回放与报告负责“讲清结果”，三者一起构成优化平台的分析闭环。",
        accent=PRIMARY,
        title_size=18,
        body_size=14,
    )
    add_footer(slide, "答辩时这一页可以顺势切换到实际页面演示，逻辑最自然。")


def slide_summary(slide) -> None:
    set_slide_bg(slide)
    add_title_block(slide, "创新点、不足与总结展望", "最后一页要既体现亮点，也保持客观，说明平台的研究价值与后续空间。")
    add_card(
        slide,
        Cm(1.2),
        Cm(3.2),
        Cm(10.2),
        Cm(10.4),
        "创新点",
        "1. 用网格化仓储拓扑替代自由路网，更贴合行走图。\n2. 将 CTU、四向车、输送线、缓存位和站台纳入同一协同闭环。\n3. 实现建模、仿真、优化、回放的一体化平台。\n4. 支持多策略实验和异常可解释分析。",
        accent=ACCENT,
        title_size=19,
        body_size=14,
    )
    add_card(
        slide,
        Cm(11.8),
        Cm(3.2),
        Cm(10.2),
        Cm(10.4),
        "当前不足",
        "1. 当前仍是本地优先 MVP，尚未对接真实 WMS/WCS/RCS。\n2. 真实数据导入与 CAD/DWG 解析尚未完成。\n3. 复杂多层资源竞争与提升机调度仍可继续增强。\n4. 学习型、自适应优化算法尚未接入。",
        accent=WARNING,
        title_size=19,
        body_size=14,
    )
    add_card(
        slide,
        Cm(22.4),
        Cm(3.2),
        Cm(10.6),
        Cm(10.4),
        "总结与展望",
        "本项目完成了一个可运行的智能仓储 AGV 调度仿真平台，实现了从拓扑建模到仿真优化再到结果分析的完整闭环。后续将继续面向真实数据接入、多层协同调度和高级优化算法推进。",
        accent=SUCCESS,
        title_size=19,
        body_size=14,
    )
    add_textbox(slide, Cm(1.2), Cm(15.1), Cm(31.8), Cm(1.0), "以上是我的汇报，感谢各位老师指导。", size=20, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)
    add_footer(slide, "结束时建议停顿一下，再进入老师提问环节。")


def write_notes() -> None:
    notes = """# 智能仓储 AGV 调度仿真优化平台 10 页答辩讲稿

## 第 1 页 封面
- 开场先说明这是一个完整的平台，不是单一算法程序。
- 平台覆盖拓扑建模、场景配置、仿真运行、算法对比和结果回放。

## 第 2 页 项目背景与问题定义
- 仓储系统中设备多、协同复杂、冲突频繁，真实系统试错成本高。
- 传统研究往往缺少系统级、可视化、可回放的实验环境。

## 第 3 页 平台目标与总体闭环
- 平台的四个目标分别是建模、仿真、优化、分析。
- 强调建拓扑到回放分析的完整闭环。

## 第 4 页 系统总体架构
- 平台分为表现层、服务层和核心层。
- 技术栈为 React + FastAPI，实时通信采用 WebSocket。

## 第 5 页 六大功能模块总览
- 依次说明驾驶舱、拓扑、场景、仿真、实验室、回放的职责。
- 突出每个页面都服务于同一个调度闭环。

## 第 6 页 拓扑编辑器与场景配置台
- 拓扑负责空间结构，场景负责运行参数。
- CTU 和四向车由拓扑设备节点自动派生到场景。

## 第 7 页 仿真工作台与协同调度闭环
- 重点讲任务不是点到点移动，而是阶段化流转。
- 说明工作台能同时看设备、任务、资源和告警。

## 第 8 页 后端调度与算法机制
- 路径规划支持 A* 和 Dijkstra，任务分配支持 Greedy 和 Hungarian。
- 平台的特点是路径、任务和资源一起耦合计算。

## 第 9 页 算法实验、结果分析与演示流程
- 算法实验室负责比较策略差异，回放页负责解释结果。
- 按驾驶舱到回放报告的顺序现场演示最清晰。

## 第 10 页 创新点、不足与总结展望
- 先讲创新点，再客观说明当前不足。
- 最后总结平台已经实现了完整闭环，并具备继续扩展到真实系统的基础。
"""
    NOTES_PATH.write_text(notes, encoding="utf-8")


def main() -> None:
    global prs
    prs = build_presentation()
    slides = [
        slide_cover,
        slide_background,
        slide_goals,
        slide_architecture,
        slide_modules,
        slide_topology_scenario,
        slide_simulation,
        slide_algorithms,
        slide_results,
        slide_summary,
    ]
    for builder in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder(slide)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)
    write_notes()
    print(f"PPT generated: {OUTPUT_PATH}")
    print(f"Notes generated: {NOTES_PATH}")


if __name__ == "__main__":
    main()
